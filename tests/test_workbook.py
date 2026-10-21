"""Offline tests for building the workbook from Reference Solutions exports."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from shi_database.naics import load_naics_targets
from shi_database.workbook import (
    ALL_SHEET,
    DEFAULT_NAICS_TABLE,
    UNMATCHED_SHEET,
    build_workbook,
    load_companies,
    sheet_name,
)

HEADER = ["Company Name", "Address", "City", "State", "ZIP Code", "IUSA Number", "Primary NAICS", "NAICS 1"]


def _write_export(path: Path, rows: list[list[str]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "20261013095301"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_build_workbook_classifies_dedupes_and_flags(tmp_path: Path) -> None:
    first = _write_export(
        tmp_path / "a.xlsx",
        [
            ["Gulf Shipyard Inc", "1 Dock Rd", "Mobile", "AL", "36602", "100", "33661101", "33661101"],
            ["Acme Engineering", "2 Main St", "Austin", "TX", "78701", "200", "54133001", "54133001"],
            ["Deep Sea Freight Co", "3 Pier St", "Miami", "FL", "33101", "300", "48311101", "48311101"],
            ["Main Street Diner", "4 Elm St", "Ames", "IA", "50010", "400", "72251101", "72251101"],
        ],
    )
    second = _write_export(
        tmp_path / "b.xlsx",
        [["Gulf Shipyard Inc", "1 Dock Rd", "Mobile", "AL", "36602", "100", "33661101", "33661101"]],
    )
    targets = load_naics_targets(DEFAULT_NAICS_TABLE)
    companies, dates, rows_read = load_companies([first, second], targets)
    assert (len(companies), rows_read, dates) == (4, 5, ["2026-10-13"])

    output = tmp_path / "out.xlsx"
    build_workbook(companies, targets, output, ["a.xlsx", "b.xlsx"], dates, rows_read)
    workbook = load_workbook(output)
    assert workbook.sheetnames[0] == "Summary"
    assert workbook[sheet_name("Shipyards")]["A2"].value == "Gulf Shipyard Inc"
    design = workbook[sheet_name("Independent Design Firms")]
    headers = [cell.value for cell in design[1]]
    assert design.cell(2, headers.index("Needs Review") + 1).value == "Independent Design Firms"
    assert design["A2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert workbook[sheet_name("Water Transportation")]["A2"].value == "Deep Sea Freight Co"
    assert workbook[UNMATCHED_SHEET]["A2"].value == "Main Street Diner"
    assert workbook[ALL_SHEET].max_row == 5
    assert workbook[ALL_SHEET].freeze_panes == "B2"


def test_load_companies_rejects_non_export(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name", "Phone"])
    workbook.save(tmp_path / "other.xlsx")
    with pytest.raises(ValueError, match="not a Reference Solutions"):
        load_companies([tmp_path / "other.xlsx"], load_naics_targets(DEFAULT_NAICS_TABLE))


def test_sheet_name_fits_excel_limit() -> None:
    assert sheet_name("Marine Engineering Organizations") == "Marine Engineering Orgs"


def test_keyword_narrowed_codes_need_keyword_in_name(tmp_path: Path) -> None:
    export = _write_export(
        tmp_path / "c.xlsx",
        [
            ["Prairie Wind Farm LLC", "1 Ridge Rd", "Ames", "IA", "50010", "500", "22111501", "22111501"],
            ["Atlantic Offshore Wind", "2 Pier Rd", "Boston", "MA", "02110", "600", "22111501", "22111501"],
            ["Coastal Hydrodynamics Inc", "3 Lab Rd", "Mobile", "AL", "36602", "700", "54133001", "54133001"],
        ],
    )
    companies, _, _ = load_companies([export], load_naics_targets(DEFAULT_NAICS_TABLE))
    by_name = {c.values["Company Name"]: c for c in companies}
    assert by_name["Prairie Wind Farm LLC"].categories == []
    assert by_name["Atlantic Offshore Wind"].categories == ["Offshore Wind"]
    hydro = by_name["Coastal Hydrodynamics Inc"]
    assert hydro.categories == ["Hydrodynamics Specialists", "Independent Design Firms"]
    assert not hydro.needs_review()


def test_keywords_match_only_at_word_start(tmp_path: Path) -> None:
    export = _write_export(
        tmp_path / "d.xlsx",
        [
            ["Acme Transport Logistics", "1 Main St", "Miami", "FL", "33101", "800", "48851001", "48851001"],
            ["Portside Freight Agents", "2 Dock St", "Miami", "FL", "33101", "900", "48851001", "48851001"],
        ],
    )
    companies, _, _ = load_companies([export], load_naics_targets(DEFAULT_NAICS_TABLE))
    by_name = {c.values["Company Name"]: c.categories for c in companies}
    assert by_name == {"Acme Transport Logistics": [], "Portside Freight Agents": ["Ports & Cargo"]}
