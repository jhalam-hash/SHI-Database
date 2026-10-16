"""Build the SHI marine-company workbook from Data Axle Reference Solutions exports.

Reference Solutions (the Data Axle web database many libraries license) exports
"Detail" spreadsheets with one row per business and up to eleven NAICS columns
holding 8-digit Data Axle codes. This module matches those codes against
``data/marine_naics_codes.csv`` and writes a formatted .xlsx.

Requires ``openpyxl`` (unlike the stdlib-only API client).
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from shi_database.naics import NaicsTarget, load_naics_targets
from shi_database.records import name_has_marine_term

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_NAICS_TABLE = REPO_ROOT / "data" / "marine_naics_codes.csv"
DEFAULT_OUTPUT = REPO_ROOT / "output" / "us_marine_companies.xlsx"

NAICS_COLUMNS = ("Primary NAICS", *(f"NAICS {i}" for i in range(1, 11)))
UNMATCHED_SHEET = "Not in SHI NAICS"
ALL_SHEET = "All Companies"

# (output header, source export column or None for a derived value)
COLUMNS: tuple[tuple[str, str | None], ...] = (
    ("Company Name", "Company Name"),
    ("Address", "Address"),
    ("City", "City"),
    ("State", "State"),
    ("ZIP Code", "ZIP Code"),
    ("County", "County"),
    ("Phone", "Phone Number Combined"),
    ("Website", "Website"),
    ("Executive", None),
    ("Executive Title", "Executive Title"),
    ("SHI Categories", None),
    ("Matched NAICS", None),
    ("Match Tier", None),
    ("Name Has Marine Term", None),
    ("Needs Review", None),
    ("Primary NAICS", None),
    ("Primary NAICS Description", "Primary NAICS Description"),
    ("All NAICS", None),
    ("Primary SIC Description", "Primary SIC Description"),
    ("Employees (Location)", "Location Employee Size Actual"),
    ("Sales Volume (Location)", "Location Sales Volume Actual"),
    ("Location Type", "Location Type"),
    ("Parent Company", "Parent Company Name"),
    ("Year Established", "Year Established"),
    ("Latitude", "Latitude"),
    ("Longitude", "Longitude"),
    ("IUSA Number", "IUSA Number"),
    ("Source File", None),
)
HEADERS = tuple(header for header, _ in COLUMNS)

_FONT = Font(name="Arial", size=10)
_BOLD = Font(name="Arial", size=10, bold=True)
_TITLE = Font(name="Arial", size=14, bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
_EXPORT_STAMP = re.compile(r"(20\d{2})(\d{2})(\d{2})\d{6}")


@dataclass
class Company:
    """One de-duplicated business with its SHI classification."""

    values: dict[str, str]
    categories: list[str]
    review: list[str]

    def needs_review(self, category: str | None = None) -> bool:
        """True if the company's place in ``category`` (or any category) is doubtful.

        That is when the category matched only through broad codes and the name has
        neither a marine term nor one of those codes' own keywords.
        """
        return bool(self.review) if category is None else category in self.review


def read_export(path: Path) -> tuple[list[dict[str, str]], str]:
    """Read the first sheet of a Reference Solutions export.

    Args:
        path: Exported .xlsx file.

    Returns:
        Rows as header-to-text mappings, and the sheet title (an export timestamp).

    Raises:
        ValueError: If the sheet lacks the columns this module relies on.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows, ())]
    missing = [col for col in ("Company Name", "Primary NAICS") if col not in header]
    if missing:
        raise ValueError(f"{path.name} is not a Reference Solutions detail export (missing {missing})")
    records = [
        {name: "" if cell is None else str(cell).strip() for name, cell in zip(header, row, strict=False)}
        for row in rows
        if any(cell not in (None, "") for cell in row)
    ]
    title = sheet.title
    workbook.close()
    return records, title


def classify(record: dict[str, str], targets: list[NaicsTarget], source: str) -> Company:
    """Match a record's NAICS codes against the SHI table and build its output row."""
    codes = [record.get(col, "")[:6] for col in NAICS_COLUMNS]
    codes = list(dict.fromkeys(code for code in codes if len(code) == 6 and code.isdigit()))
    name = record.get("Company Name", "")
    # A code narrowed by its own keywords (e.g. 221115 wind power, "offshore") only counts when
    # the name contains one, mirroring the keyword queries the API pull runs for it.
    hits = [
        target
        for target in targets
        if target.code in codes
        and (not target.keywords or any(keyword.lower() in name.lower() for keyword in target.keywords))
    ]
    categories = sorted({target.category for target in hits})
    has_core = any(not target.is_broad for target in hits)
    marine_name = name_has_marine_term(name)
    review = []
    for category in categories:
        category_hits = [target for target in hits if target.category == category]
        if marine_name or any(not target.is_broad or target.keywords for target in category_hits):
            continue
        review.append(category)
    tier = "core" if has_core else ("broad" if hits else "")

    values = {header: record.get(source_col, "") if source_col else "" for header, source_col in COLUMNS}
    executive = f"{record.get('Executive First Name', '')} {record.get('Executive Last Name', '')}"
    values.update(
        {
            "Executive": executive.strip(),
            "SHI Categories": "; ".join(categories),
            "Matched NAICS": "; ".join(sorted({target.code for target in hits})),
            "Match Tier": tier,
            "Name Has Marine Term": "yes" if marine_name else "no",
            "Needs Review": "; ".join(review),
            "Primary NAICS": codes[0] if codes else "",
            "All NAICS": "; ".join(codes),
            "Source File": source,
        }
    )
    return Company(values, categories, review)


def dedupe_key(values: dict[str, str]) -> str:
    """Key on Data Axle's IUSA number, falling back to name, address and ZIP."""
    if values["IUSA Number"]:
        return values["IUSA Number"]
    return "|".join((values["Company Name"], values["Address"], values["ZIP Code"])).lower()


def load_companies(paths: Iterable[Path], targets: list[NaicsTarget]) -> tuple[list[Company], list[str], int]:
    """Read, classify and de-duplicate every export.

    Returns:
        Unique companies sorted by state, city and name; the export dates found
        in sheet titles (YYYY-MM-DD); and the number of input rows read.
    """
    companies: dict[str, Company] = {}
    export_dates: set[str] = set()
    total = 0
    for path in paths:
        records, title = read_export(path)
        total += len(records)
        stamp = _EXPORT_STAMP.fullmatch(title)
        if stamp:
            export_dates.add("-".join(stamp.groups()))
        for record in records:
            company = classify(record, targets, path.name)
            companies.setdefault(dedupe_key(company.values), company)
    ordered = sorted(
        companies.values(),
        key=lambda c: (c.values["State"], c.values["City"], c.values["Company Name"].lower()),
    )
    return ordered, sorted(export_dates), total


def sheet_name(category: str) -> str:
    """Shorten a category to Excel's 31-character sheet-name limit."""
    return category.replace("Organizations", "Orgs")[:31]


def write_company_sheet(
    sheet: Worksheet, companies: list[Company], review_category: str | None, flag_review: bool = True
) -> None:
    """Write a company table with bold frozen headers, filters and review highlighting.

    Rows are highlighted when the company needs review for ``review_category``, or
    for any category when it is None.
    """
    sheet.append(list(HEADERS))
    for company in companies:
        sheet.append([company.values[header] for header in HEADERS])
        if flag_review and company.needs_review(review_category):
            for cell in sheet[sheet.max_row]:
                cell.fill = _REVIEW_FILL
    for cell in sheet[1]:
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = _FONT
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 2)}"
    fit_columns(sheet)


def fit_columns(sheet: Worksheet, minimum: int = 8, maximum: int = 45) -> None:
    """Size columns to their longest value within sensible bounds."""
    widths: Counter[int] = Counter()
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and not str(cell.value).startswith("="):
                widths[cell.column] = max(widths[cell.column], len(str(cell.value)))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, minimum), maximum)


def write_summary(
    sheet: Worksheet,
    companies: list[Company],
    targets: list[NaicsTarget],
    sources: list[str],
    export_dates: list[str],
    rows_read: int,
) -> None:
    """Write counts (as live formulas over the All Companies tab) and a method note."""
    all_ref = f"'{ALL_SHEET}'"
    col = {header: get_column_letter(i) for i, header in enumerate(HEADERS, start=1)}
    last = len(companies) + 1

    def rng(header: str) -> str:
        return f"{all_ref}!${col[header]}$2:${col[header]}${last}"

    def section(title: str, columns: tuple[str, ...]) -> None:
        sheet.append([])
        sheet.append([title])
        sheet.cell(sheet.max_row, 1).font = _TITLE
        sheet.append(list(columns))
        for cell in sheet[sheet.max_row]:
            cell.font = _BOLD
            cell.fill = _HEADER_FILL

    sheet.append(["SHI U.S. Marine Companies"])
    sheet["A1"].font = Font(name="Arial", size=16, bold=True)
    pulled = ", ".join(export_dates) or "unknown"
    notes = [
        ("Data pulled (export date)", pulled),
        ("Workbook built", date.today().isoformat()),
        ("Source", "Data Axle Reference Solutions detail exports: " + ", ".join(sources)),
        ("Rows read / unique companies", f"{rows_read} / {len(companies)} (de-duplicated on IUSA Number)"),
        (
            "Method",
            "Each company's Primary NAICS and NAICS 1-10 (8-digit Data Axle codes, first 6 digits used) "
            "are matched against data/marine_naics_codes.csv. A company appears on every category tab "
            "its codes match. 'core' codes are marine by definition; 'broad' codes (e.g. 541330 "
            "Engineering Services) are not. A broad code with its own keywords (e.g. 221115 wind power: "
            "'offshore') counts only when the company name contains one. When a category matched only "
            "through other broad codes and the name has no marine term, the category is listed under "
            "Needs Review and the row is highlighted yellow on that tab and on All Companies. "
            "Companies with no SHI code are on the "
            f"'{UNMATCHED_SHEET}' tab.",
        ),
        (
            "Coverage caveat",
            "Counts reflect only the searches that were exported, not every U.S. company under "
            "these codes. Run Reference Solutions searches on the SHI NAICS codes to fill gaps.",
        ),
    ]
    for label, value in notes:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = _BOLD
        sheet.cell(sheet.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    section("Companies by SHI category", ("Category", "Companies", "Needs Review"))
    categories = list(dict.fromkeys(target.category for target in targets))
    for category in categories:
        tab = f"'{sheet_name(category)}'"
        sheet.append(
            [
                category,
                f"=COUNTA({tab}!$A:$A)-1",
                f'=COUNTIF({tab}!${col["Needs Review"]}:${col["Needs Review"]},"*{category}*")',
            ]
        )
    sheet.append(["No SHI NAICS match", f"=COUNTA('{UNMATCHED_SHEET}'!$A:$A)-1", ""])
    sheet.append(["All unique companies", f"=COUNTA({all_ref}!$A:$A)-1", ""])
    sheet.cell(sheet.max_row, 1).font = _BOLD

    section("Companies by SHI NAICS code (any listed code)", ("NAICS", "Title", "Tier", "Companies"))
    for target in targets:
        sheet.append(
            [target.code, target.title, target.tier, f'=COUNTIF({rng("Matched NAICS")},"*{target.code}*")']
        )

    section("Companies by primary NAICS (as exported)", ("Primary NAICS", "Description", "Companies"))
    descriptions = {c.values["Primary NAICS"]: c.values["Primary NAICS Description"] for c in companies}
    primary_counts = Counter(c.values["Primary NAICS"] for c in companies)
    for code, _ in primary_counts.most_common():
        sheet.append([code, descriptions[code], f'=COUNTIF({rng("Primary NAICS")},"{code}")'])

    section("Companies by state", ("State", "Companies"))
    for state in sorted({c.values["State"] for c in companies}):
        sheet.append([state, f'=COUNTIF({rng("State")},"{state}")'])

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if not cell.font.bold:
                cell.font = _FONT
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 12


def build_workbook(
    companies: list[Company],
    targets: list[NaicsTarget],
    output: Path,
    sources: list[str],
    export_dates: list[str],
    rows_read: int,
) -> None:
    """Write the Summary, per-category, unmatched and All Companies tabs."""
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None  # noqa: S101 - a new Workbook always has an active sheet
    summary.title = "Summary"
    broad_only = {t.category for t in targets} - {t.category for t in targets if not t.is_broad}
    for category in dict.fromkeys(target.category for target in targets):
        members = [c for c in companies if category in c.categories]
        write_company_sheet(workbook.create_sheet(sheet_name(category)), members, category)
        if category in broad_only:
            workbook[sheet_name(category)].sheet_properties.tabColor = "BF8F00"
    unmatched = [c for c in companies if not c.categories]
    write_company_sheet(workbook.create_sheet(UNMATCHED_SHEET), unmatched, None, flag_review=False)
    write_company_sheet(workbook.create_sheet(ALL_SHEET), companies, None)
    write_summary(summary, companies, targets, sources, export_dates, rows_read)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main(argv: list[str] | None = None) -> int:
    """Build the workbook from one or more Reference Solutions exports."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("exports", nargs="+", type=Path, help="Reference Solutions detail .xlsx files")
    parser.add_argument("--naics-table", type=Path, default=DEFAULT_NAICS_TABLE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    try:
        targets = load_naics_targets(args.naics_table)
        companies, export_dates, rows_read = load_companies(args.exports, targets)
    except (FileNotFoundError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    sources = [path.name for path in args.exports]
    build_workbook(companies, targets, args.output, sources, export_dates, rows_read)
    matched = sum(1 for c in companies if c.categories)
    print(
        f"wrote {len(companies)} unique companies ({matched} in SHI categories) to {args.output}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
